import os
import streamlit as st
import pandas as pd
from datetime import date, datetime
from io import BytesIO
from docx import Document

def gen_papeleta_vac(apellidos, nombres, dni_b, position, f_ingreso, period, start_d, end_d, days):
    if not os.path.exists("Template_Papeleta.docx"):
        return st.error("⚠️ No se encontró la plantilla 'Template_Papeleta.docx'.")

    doc = Document("Template_Papeleta.docx")
    fin_dt = pd.to_datetime(end_d, errors="coerce")
    retorno_dt = (fin_dt + pd.Timedelta(days=1 if fin_dt.weekday() < 5 else (2 if fin_dt.weekday() == 5 else 1))) if pd.notna(fin_dt) else None
  
    hoy = date.today()
    meses = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]
  
    reps = {
        "{{APELLIDOS}}": str(apellidos).upper(), "{{NOMBRES}}": str(nombres).upper(),
        "{{DNI}}": str(dni_b), "{{CARGO}}": str(position).upper(),
        "{{F_INGRESO}}": f_ingreso.strftime("%d/%m/%Y") if isinstance(f_ingreso, (date, datetime)) else str(f_ingreso),
        "{{PERIODO}}": str(period), "{{DIAS}}": str(days),
        "{{F_INICIO}}": start_d.strftime("%d/%m/%Y") if isinstance(start_d, (date, datetime)) else str(start_d),
        "{{F_FIN}}": end_d.strftime("%d/%m/%Y") if isinstance(end_d, (date, datetime)) else str(end_d),
        "{{F_RETORNO}}": retorno_dt.strftime("%d/%m/%Y") if retorno_dt else "",
        "{{FECHA_FIRMA}}": f"Huancayo, {hoy.day} de {meses[hoy.month-1]} de {hoy.year}"
    }

    def replace_text(element):
        for run in element.runs:
            for k, v in reps.items():
                if k in run.text: run.text = run.text.replace(k, v)

    for p in doc.paragraphs: replace_text(p)
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                for p in c.paragraphs: replace_text(p)

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
